"""
Builds the four free biodata templates (boy/girl x Word/PDF) into public/downloads/.
Run from the project root:  python scripts/build-biodata-templates.py
Requires: pip install python-docx reportlab
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

OUT = Path(__file__).resolve().parent.parent / "public" / "downloads"
OUT.mkdir(parents=True, exist_ok=True)
FOOTER = "Free template from HumNikah – humnikah.com"

# (section, [(label, placeholder)])
BOY = [
    ("", [
        ("Name", "[Full name or initials]"),
        ("Date of birth / Age", "[DD/MM/YYYY] / [Age]"),
        ("Height", "[ft in / cm]"),
        ("Marital status", "[Never married / Divorced / Widower]"),
        ("Languages", "[e.g. Urdu, English, Konkani, Hindi]"),
        ("School of thought", "[Optional]"),
    ]),
    ("EDUCATION AND WORK", [
        ("Education", "[Degree, institution]"),
        ("Profession", "[Job title / business]"),
        ("City of work", "[City]"),
    ]),
    ("RELIGIOUS PRACTICE", [
        ("Practice", "[e.g. prays five times a day, keeps beard - optional]"),
    ]),
    ("FAMILY DETAILS", [
        ("Father", "[Name, occupation]"),
        ("Mother", "[Name, occupation]"),
        ("Siblings", "[Number, brothers / sisters, married or unmarried]"),
        ("Family home", "[City / area]"),
    ]),
    ("EXPECTATIONS", [
        ("Age range", "[Range]"),
        ("Education", "[Preference]"),
        ("Location", "[Preference]"),
        ("Other", "[Family values, religious practice, etc.]"),
    ]),
    ("CONTACT", [
        ("Contact through", "[Father / guardian / HumNikah]"),
        ("Photograph", "[Available on request]"),
    ]),
]

GIRL = [
    ("", [
        ("Name", "[Full name or initials]"),
        ("Date of birth / Age", "[DD/MM/YYYY] / [Age]"),
        ("Height", "[ft in / cm]"),
        ("Marital status", "[Never married / Divorced / Widow]"),
        ("Languages", "[e.g. Urdu, English, Tamil, Hindi]"),
        ("School of thought", "[Optional]"),
    ]),
    ("EDUCATION AND WORK", [
        ("Education", "[Degree, institution]"),
        ("Profession", "[Job title / homemaker / studying]"),
        ("City", "[City]"),
        ("Plans after marriage", "[e.g. continue working / open to discussion]"),
    ]),
    ("RELIGIOUS PRACTICE", [
        ("Practice", "[e.g. prays five times a day, hijab - optional]"),
    ]),
    ("FAMILY DETAILS", [
        ("Father", "[Name, occupation]"),
        ("Mother", "[Name, occupation]"),
        ("Siblings", "[Number, brothers / sisters, married or unmarried]"),
        ("Wali", "[Name, relation]"),
        ("Family home", "[City / area]"),
    ]),
    ("EXPECTATIONS", [
        ("Age range", "[Range]"),
        ("Education", "[Preference]"),
        ("Location", "[Preference]"),
        ("Other", "[Family values, religious practice, etc.]"),
    ]),
    ("CONTACT", [
        ("Contact through", "[Wali / guardian / HumNikah]"),
        ("Photograph", "[Available on request]"),
    ]),
]

TEMPLATES = {"boy": BOY, "girl": GIRL}


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def build_docx(kind: str, rows):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.4)
    doc.core_properties.title = f"Muslim Biodata Format for Marriage – {kind.capitalize()}"
    doc.core_properties.author = "HumNikah"

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("BIODATA FOR MARRIAGE")
    run.bold = True
    run.font.size = Pt(16)

    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for heading, fields in rows:
        if heading:
            cells = table.add_row().cells
            merged = cells[0].merge(cells[1])
            merged.text = ""
            r = merged.paragraphs[0].add_run(heading)
            r.bold = True
            shade(merged, "EFE7DA")
        for label, hint in fields:
            cells = table.add_row().cells
            cells[0].text = ""
            cells[0].paragraphs[0].add_run(label).bold = True
            cells[1].text = hint
            cells[0].width, cells[1].width = Cm(5), Cm(12.4)

    foot = doc.add_paragraph()
    foot.paragraph_format.space_before = Pt(10)
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = foot.add_run(FOOTER)
    fr.italic = True
    fr.font.size = Pt(9)
    doc.save(OUT / f"muslim-biodata-format-{kind}.docx")


def build_pdf(kind: str, rows):
    path = OUT / f"muslim-biodata-format-{kind}.pdf"
    pdf = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.4 * cm, bottomMargin=1.4 * cm,
        title=f"Muslim Biodata Format for Marriage – {kind.capitalize()}", author="HumNikah",
    )
    title_style = ParagraphStyle("t", fontName="Helvetica-Bold", fontSize=16, alignment=1, spaceAfter=8)
    cell = ParagraphStyle("c", fontName="Helvetica", fontSize=10, leading=13)
    bold = ParagraphStyle("b", parent=cell, fontName="Helvetica-Bold")
    foot = ParagraphStyle("f", fontName="Helvetica-Oblique", fontSize=8.5, alignment=1, spaceBefore=10)

    data, styles = [], [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b9ae9b")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for heading, fields in rows:
        if heading:
            idx = len(data)
            data.append([Paragraph(heading, bold), ""])
            styles += [("SPAN", (0, idx), (1, idx)), ("BACKGROUND", (0, idx), (1, idx), colors.HexColor("#EFE7DA"))]
        for label, hint in fields:
            data.append([Paragraph(label, bold), Paragraph(hint, cell)])

    table = Table(data, colWidths=[5 * cm, 12.4 * cm])
    table.setStyle(TableStyle(styles))
    pdf.build([Paragraph("BIODATA FOR MARRIAGE", title_style), table, Spacer(1, 6), Paragraph(FOOTER, foot)])


for kind, rows in TEMPLATES.items():
    build_docx(kind, rows)
    build_pdf(kind, rows)
print("built:", sorted(p.name for p in OUT.iterdir()))
